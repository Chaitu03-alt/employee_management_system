"""OpenPyXL workbook generation for employee reports."""

from __future__ import annotations

from collections import defaultdict
from collections.abc import Iterable
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from employee import EMPLOYEE_FIELDS, Employee


class ExcelReportGenerator:
    """Create a formatted Excel workbook containing employee reports."""

    HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
    HEADER_FONT = Font(color="FFFFFF", bold=True)
    CURRENCY_FORMAT = '"INR" #,##0'
    DECIMAL_FORMAT = "0.00"

    def __init__(self, employees: Iterable[Employee], reports_directory: str | Path) -> None:
        self.employees = list(employees)
        self.reports_directory = Path(reports_directory)

    def generate(self, filename: str = "employee_reports.xlsx") -> Path:
        """Generate the four-sheet employee report workbook."""
        self.reports_directory.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        workbook.remove(workbook.active)

        self._create_employee_report(workbook)
        self._create_department_report(workbook)
        self._create_salary_report(workbook)
        self._create_performance_report(workbook)

        output_path = self.reports_directory / filename
        workbook.save(output_path)
        return output_path

    def _create_employee_report(self, workbook: Workbook) -> None:
        """Add a detailed employee-record worksheet."""
        sheet = workbook.create_sheet("Employee Report")
        sheet.append(EMPLOYEE_FIELDS)
        for employee in self.employees:
            salary = self._number(employee.salary)
            experience = self._number(employee.experience)
            rating = self._number(employee.performance_rating)
            sheet.append(
                [
                    employee.employee_id,
                    employee.name,
                    employee.department,
                    salary if salary is not None else employee.salary,
                    experience if experience is not None else employee.experience,
                    employee.city,
                    employee.joining_date,
                    rating if rating is not None else employee.performance_rating,
                ]
            )
        self._style_sheet(sheet, (14, 26, 18, 16, 14, 16, 16, 21))
        for cell in sheet["D"][1:]:
            cell.number_format = self.CURRENCY_FORMAT

    def _create_department_report(self, workbook: Workbook) -> None:
        """Add a department-level employee and compensation summary."""
        sheet = workbook.create_sheet("Department Report")
        sheet.append(["Department", "Employee Count", "Average Salary", "Average Experience"])
        groups: dict[str, list[Employee]] = defaultdict(list)
        for employee in self.employees:
            groups[employee.department or "Unknown"].append(employee)

        for department in sorted(groups, key=str.casefold):
            employees = groups[department]
            salaries = self._numbers(employee.salary for employee in employees)
            experience = self._numbers(employee.experience for employee in employees)
            sheet.append(
                [
                    department,
                    len(employees),
                    self._average(salaries),
                    self._average(experience),
                ]
            )
        self._style_sheet(sheet, (20, 18, 18, 22))
        for cell in sheet["C"][1:]:
            cell.number_format = self.CURRENCY_FORMAT
        for cell in sheet["D"][1:]:
            cell.number_format = self.DECIMAL_FORMAT

    def _create_salary_report(self, workbook: Workbook) -> None:
        """Add salary statistics and distribution bands."""
        sheet = workbook.create_sheet("Salary Report")
        salaries = self._numbers(employee.salary for employee in self.employees)
        sheet.append(["Salary Summary", "Value"])
        sheet.append(["Employee Count", len(salaries)])
        sheet.append(["Average Salary", self._average(salaries)])
        sheet.append(["Minimum Salary", min(salaries) if salaries else None])
        sheet.append(["Maximum Salary", max(salaries) if salaries else None])
        sheet.append([])
        sheet.append(["Salary Range", "Employee Count"])

        salary_bands = (
            ("Below INR 500,000", 0, 500_000),
            ("INR 500,000 - INR 749,999", 500_000, 750_000),
            ("INR 750,000 - INR 999,999", 750_000, 1_000_000),
            ("INR 1,000,000 - INR 1,499,999", 1_000_000, 1_500_000),
            ("INR 1,500,000 and above", 1_500_000, None),
        )
        for label, minimum, maximum in salary_bands:
            count = sum(salary >= minimum and (maximum is None or salary < maximum) for salary in salaries)
            sheet.append([label, count])

        self._style_sheet(sheet, (28, 20), header_rows=(1, 7))
        for row in range(3, 6):
            sheet.cell(row=row, column=2).number_format = self.CURRENCY_FORMAT

    def _create_performance_report(self, workbook: Workbook) -> None:
        """Add performance-rating counts, percentages, and average salaries."""
        sheet = workbook.create_sheet("Performance Report")
        sheet.append(["Performance Rating", "Employee Count", "Percentage", "Average Salary"])
        groups: dict[str, list[Employee]] = defaultdict(list)
        for employee in self.employees:
            groups[employee.performance_rating or "Unknown"].append(employee)

        total = len(self.employees)
        for rating in sorted(groups, key=self._rating_sort_key):
            employees = groups[rating]
            salaries = self._numbers(employee.salary for employee in employees)
            sheet.append([rating, len(employees), len(employees) / total if total else 0, self._average(salaries)])
        self._style_sheet(sheet, (24, 18, 14, 18))
        for cell in sheet["C"][1:]:
            cell.number_format = "0.00%"
        for cell in sheet["D"][1:]:
            cell.number_format = self.CURRENCY_FORMAT

    def _style_sheet(
        self, sheet: Worksheet, widths: tuple[int, ...], header_rows: tuple[int, ...] = (1,)
    ) -> None:
        """Apply reusable table formatting to a report worksheet."""
        for header_row in header_rows:
            for cell in sheet[header_row]:
                cell.fill = self.HEADER_FILL
                cell.font = self.HEADER_FONT
                cell.alignment = Alignment(horizontal="center")
        for column_index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + column_index)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    @staticmethod
    def _number(value: str) -> float | None:
        """Convert salary-like or numeric text into a number when possible."""
        try:
            return float(str(value).replace(",", "").replace("INR", "").strip())
        except (TypeError, ValueError):
            return None

    @classmethod
    def _numbers(cls, values: Iterable[str]) -> list[float]:
        """Return the valid numeric values from an iterable."""
        return [number for value in values if (number := cls._number(value)) is not None]

    @staticmethod
    def _average(values: list[float]) -> float | None:
        """Calculate an average while safely handling an empty collection."""
        return sum(values) / len(values) if values else None

    @staticmethod
    def _rating_sort_key(rating: str) -> tuple[int, float | str]:
        """Order numeric ratings before non-numeric ratings."""
        try:
            return (0, float(rating))
        except ValueError:
            return (1, rating.casefold())
